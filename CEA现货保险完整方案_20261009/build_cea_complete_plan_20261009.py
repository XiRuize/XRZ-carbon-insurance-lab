"""Create the complete Chinese proposal with editable native Word equations."""
from pathlib import Path
from copy import deepcopy
import csv,json,zipfile
from docx import Document
from docx.shared import Inches,Pt,Mm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
OUT=ROOT if ROOT.name=='CEA现货保险完整方案_20261009' else ROOT/'CEA现货保险完整方案_20261009'
S=json.loads((OUT/'统一计算结果.json').read_text(encoding='utf-8'))
E=S['enterprise'];P=S['contract'];J=S['joint']
def csvrows(name):
    with (OUT/'data'/name).open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
metrics=csvrows('保费算法回放指标.csv');candidates=csvrows('全部采购投保候选.csv');sales=csvrows('富余配额出售比较.csv')
daily=[r for r in csvrows('CEA输入日行情.csv') if r['date'].startswith('2026')]
def money(x):return f'{float(x):,.2f}'
def price(x):return f'{float(x):.2f}'
def percent(x):return f'{float(x)*100:.2f}%'

def el(tag,*children,**attrs):
    x=OxmlElement('m:'+tag)
    for k,v in attrs.items():x.set(qn('m:'+k),str(v))
    for c in children:x.append(deepcopy(c) if not isinstance(c,str) else mr(c))
    return x
def mr(text,plain=False):
    x=el('r')
    if plain:x.append(el('rPr',el('sty',val='p')))
    t=el('t');t.text=str(text);t.set('{http://www.w3.org/XML/1998/namespace}space','preserve');x.append(t);return x
def seq(*parts):return [mr(p) if isinstance(p,(str,int,float)) else p for p in parts]
def wrap(tag,p):return el(tag,*(p if isinstance(p,list) else [p]))
def sub(b,i):return el('sSub',wrap('e',b if isinstance(b,list) else seq(b)),wrap('sub',i if isinstance(i,list) else seq(i)))
def sup(b,i):return el('sSup',wrap('e',b if isinstance(b,list) else seq(b)),wrap('sup',i if isinstance(i,list) else seq(i)))
def frac(a,b):return el('f',wrap('num',a if isinstance(a,list) else seq(a)),wrap('den',b if isinstance(b,list) else seq(b)))
def group(p,left='(',right=')'):return el('d',el('dPr',el('begChr',val=left),el('endChr',val=right)),wrap('e',p if isinstance(p,list) else seq(p)))
def fn(n,p):return el('func',wrap('fName',seq(mr(n,True))),wrap('e',group(p)))
def summ(i,n,p):return el('nary',el('naryPr',el('chr',val='∑'),el('limLoc',val='undOvr')),wrap('sub',i if isinstance(i,list) else seq(i)),wrap('sup',seq(n)),wrap('e',p if isinstance(p,list) else seq(p)))
def root(p):return el('rad',el('radPr',el('degHide',val=1)),el('deg'),wrap('e',p if isinstance(p,list) else seq(p)))

doc=Document();sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(18);sec.bottom_margin=Mm(18);sec.left_margin=Mm(18);sec.right_margin=Mm(18)
sec.footer_distance=Mm(8)
for name,size,east,bold in [('Normal',11,'宋体',False),('Title',22,'黑体',True),('Subtitle',12,'黑体',False),('Heading 1',16,'黑体',True),('Heading 2',12.5,'黑体',True),('Heading 3',11.5,'黑体',True)]:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.font.bold=bold
    pr=st.element.get_or_add_rPr();rf=pr.find(qn('w:rFonts'))
    if rf is None:rf=OxmlElement('w:rFonts');pr.append(rf)
    rf.set(qn('w:eastAsia'),east)
for b in list(doc.styles.element.iter(qn('w:pBdr'))):b.getparent().remove(b)
pf=doc.styles['Normal'].paragraph_format;pf.line_spacing=1.12;pf.space_after=Pt(5)
for name in ['Heading 1','Heading 2','Heading 3']:
    pf=doc.styles[name].paragraph_format;pf.keep_with_next=True;pf.space_before=Pt(8);pf.space_after=Pt(6)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('第 ');fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld);footer.add_run(' 页')
for r in footer.runs:r.font.size=Pt(9)
md=[];equations=[]
def para(text,bold=False):
    p=doc.add_paragraph();p.paragraph_format.widow_control=True;r=p.add_run(text);r.bold=bold
    md.append(text+'\n');return p
def heading(text,level=2):doc.add_heading(text,level);md.append('#'*(level+1)+' '+text+'\n')
def page(text):doc.add_page_break();heading(text,1)
def eq(parts,name,tex):
    n=len(equations)+1;equations.append(dict(number=n,name=name,latex=tex))
    p=doc.add_paragraph();p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(7)
    p.paragraph_format.left_indent=Inches(.08);p.paragraph_format.tab_stops.add_tab_stop(Inches(6.65),WD_TAB_ALIGNMENT.RIGHT)
    p._p.append(el('oMath',*parts));r=p.add_run('\t'+f'（{n}）');r.font.size=Pt(9)
    md.append('\\[\n'+tex+'\n\\]\n');return p
def table(headers,rows,widths=None,size=10):
    if widths is None:widths=[6.85/len(headers)]*len(headers)
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,h in enumerate(headers):t.columns[i].width=Inches(widths[i]);t.rows[0].cells[i].text=str(h)
    for row in rows:
        c=t.add_row().cells
        for i,v in enumerate(row):c[i].text=str(v)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if ri==0:pr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for side in ['top','bottom','left','right']:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1' if ri==0 else ('F6F7F9' if ri%2==0 else 'FFFFFF'));cp.append(sh)
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 and len(headers)<6 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before=Pt(1);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.04
                for r in p.runs:r.font.size=Pt(size);r.bold=(ri==0)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    md.append('| '+' | '.join(map(str,headers))+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+''.join('| '+' | '.join(str(x).replace('\n',' ') for x in row)+' |\n' for row in rows)+'\n')
    return t
def steps(lines):
    for i,text in enumerate(lines,1):para(f'{i}. {text}')
def code(text):
    for line in text.splitlines():
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.03
        r=p.add_run(line);r.font.name='Consolas';r.font.size=Pt(9.5);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    md.append('```text\n'+text+'\n```\n')

doc.add_paragraph('CEA现货采购保险及企业购售决策完整方案',style='Title')
doc.add_paragraph('模型建立 产品定价 承保限额 企业方案与智能体交互',style='Subtitle')
md.append('# CEA现货采购保险及企业购售决策完整方案\n')
para('研究日期 2026年10月9日\n行情截止 2026年9月30日\n面向 保险创新研究团队与企业碳资产管理人员')
heading('方案摘要')
para('本方案建立以CEA现货指数为依据的企业采购成本上涨保险。企业先申报真实未锁价采购缺口，再选择保障阶段、承保吨数、免赔价差、赔付比例和最高保障价差。智能体同时比较提前采购、部分投保与富余配额出售方案，帮助企业管理净成本、资金占用和不利价格情景。')
para('本版统一采用最近60个收益率、5日循环区块重采样和100,000条价格路径，报价与企业决策共用同一组情景。候选条款为末5交易日均价结算、免赔max(2元/吨,2%预算价)、80%赔付以及预算价5%的最高保障价差。')
para(f'每档1,000吨时，10、20、40、60交易日示范保费分别为{money(S["quotes"][0]["gross_premium_cny"])}、{money(S["quotes"][1]["gross_premium_cny"])}、{money(S["quotes"][2]["gross_premium_cny"])}、{money(S["quotes"][3]["gross_premium_cny"])}元，每档最高赔款3,798.64元。400、300、300吨分三阶段保障的示范总保费为{money(J["gross_premium_cny"])}元。',True)
para('历史回放发现基准模型在后段低估赔款，且企业实际资料与商业费用参数尚待核验。文中数值用于研究比较，不形成正式保险报价或当前市场交易指令。文档附181条实际使用的日行情，完整来源、计算结果和程序随数据包提供。')
heading('阅读导航')
table(['章节','主要内容'],[['1至3','业务目标 数据与符号'],['4至5','价格模型与预测结果'],['6至10','保险责任 限额 定价及多阶段产品'],['11至12','历史回放与降低保费的条件'],['13至18','企业采购 投保与富余配额出售'],['19至21','智能体交互 期货辅助与实施顺序'],['附录','181条日行情 算法步骤与引用来源']],[1.15,5.7])

page('1 业务目标与产品结构')
heading('1 1 企业采购保险的目标')
para('企业年度排放量并不等于需要购买的配额量。产品保障尚未锁定价格的实际采购需求，赔款随约定现货指数上涨而增加，并受免赔、赔付比例和责任限额约束。企业在采购前安排保障；保障数量及利益在风险发生时仍需满足最终合同要求。')
heading('1 2 产品结构')
table(['要素','候选设计'],[['保险对象','有经核验的未来CEA现货采购需求的企业'],['保险责任','约定阶段现货指数超过预算价与免赔价差后按公式赔付'],['期限','10 20 40 60交易日单阶段 或多个采购阶段的主保单'],['指数','每阶段预先约定的末5个交易日收盘均价'],['数量','未锁价净需求 本次申请量 保险人批准容量共同约束'],['金额','承保数量 × 每吨最高保障价差 × 赔付比例'],['企业决策','采购量和日期 部分投保量 富余配额出售条件']],[1.3,5.55])
heading('1 3 经济效果的评价')
para('采购方案比较扣除赔款后的净支出，加入保费、交易费用及资金成本。出售方案比较净回款；需要取得成本等资料才能计算利润。保险可能提高平均净支出，同时降低不利情景中的成本，因此是否值得购买还取决于预算约束、现金紧张及经营损失。')
para('企业可以买卖配额，同时承担足额清缴义务；全国交易规则对账户可交易量和可用资金设有约束。[S1][S2] 本方案在出售前预留履约量和安全缓冲，并检查配额年份、交付与可结转条件。')
heading('1 4 方案的成熟度')
para('价格情景、赔款函数、单阶段和多阶段数值求解、数量校验及企业策略枚举已形成计算原型。智能体界面、证据核验、统一数量台账及商业承保流程仍需接入。长期年度成本指数保险需要独立的数据与期限校准；本版数值聚焦短期CEA采购。')

page('2 数据来源与参数')
table(['数据项目','本版数值或口径'],[['研究使用行情','2026年1月5日至9月30日 连续181条'],['来源结构','149条官方日行情 32条转载行情'],['保留原始行情','共'+str(S['rows_all'])+'条 含日期 来源网址与质量标记'],['最后单日收盘','94.96元/吨'],['2026年样本最低及最高','72.50元/吨 99.69元/吨'],['最近60日对数收益均值',f'{S["daily_log_mean60"]:.8f} 每交易日'],['最近60日收益标准差',f'{S["daily_log_sd60"]:.8f} 每交易日'],['模拟配置','60收益率 5日区块 100,000条60日路径'],['统一随机种子','2026100903 共享路径用于预测 定价与企业方案']],[2.0,4.85])
heading('2 1 预算价所用的五条数据')
table(['日期','收盘价 元每吨'],[[x['date'],price(x['close'])] for x in S['last5_closes']],[3.45,3.4])
para('五条价格均为研究起点之前已知的交易日数据，算术均值为94.966元/吨。9月25日及国庆等休市日不作交易日收益计算。181条价格附于附录，来源网址和质量标记保留于数据文件。')
heading('2 2 使用与核验要求')
para('只使用连续2026年样本进行当前模拟，避免把早期长期缺口当作每日收益。转载记录尚非全部逐日完成官方交叉验证，正式报价需刷新至报价时点并核验数据。交易日历沿用已有2026年安排，换起点或跨年份时重新检查。')
para('本报告研究日期为10月9日，报价起点仍为9月30日，不含10月8日或9日的新行情。价格区间是模型分位范围，不是对真实覆盖概率已经校准的承诺。')

page('3 符号与单位')
table(['符号','含义','单位'],[['Sₜ P实际','现货指数 实际执行价','元/吨'],['h Tᵢ τᵢ','交易日期限 阶段日期 日历年数','交易日 日期 年'],['R H C B','履约需求 可用资源 锁价采购 已分配需求','吨'],['D q申请 q容量','净未锁价需求 申请量 保险人容量','吨'],['Q Qᵢ Q保险','承保数量 阶段数量 企业投保量','吨'],['K d c η','预算价 免赔价差 最高价差 赔付比例','元/吨 元/吨 元/吨 比例'],['Xᵢ Lᵢ Aᵢ','阶段指数 阶段赔款 最高赔款','元/吨 元 元'],['rⱼ zⱼ w b','对数收益 去均值收益 窗口 区块长度','无量纲 无量纲 60 5'],['M μ TVaRα','路径数 期望赔款 上尾平均赔款','100,000 元 元'],['r k v p F','折现率 资本成本率 费用率 利润率 固定费用','年率 年率 比例 比例 元'],['Π纯 Π EC','折现纯保费 总保费 资本代理量','元'],['q₀ θ α','提前买入量 总缺口覆盖率 剩余采购覆盖率','吨 比例 比例'],['I bₜ sₜ R安全','库存 买入量 卖出量 履约及缓冲预留','吨'],['C净 N λ','企业净成本 出售净回款 风险权重','元 元 无量纲']],[1.6,3.95,1.3])
para('c表示免赔后最高纳入赔付的每吨价差。赔款上沿为K+d+c，不能将c=5%K理解为所有高于预算价5%的涨幅都已覆盖。α在企业决策中表示数量比例；尾部概率水平以具体95%或99.5%标注，避免混淆。')

page('4 CEA价格情景模型')
heading('4 1 收益率与中心化')
eq(seq(sub('r','j'),'=',fn('ln',frac(sub('S','j'),sub('S','j−1')))),'日对数收益',r'r_j=\log(S_j/S_{j-1})')
eq(seq(sub('r','mean'),'=',frac('1','60'),summ('j=1','60',sub('r','j')),',  ',sub('z','j'),'=',sub('r','j'),'−',sub('r','mean')),'收益中心化',r'\bar r=\frac1{60}\sum_{j=1}^{60}r_j,\quad z_j=r_j-\bar r')
para('每个起点只读取当时已知的61个收盘价。去均值移除窗口平均对数漂移，保留短期波动。对数漂移为零时，价格算术均值仍可能因指数变换而略高于起点价；该过程不等于市场风险中性定价。')
heading('4 2 循环区块重采样')
eq(seq(sup(sub('r',seq('m,t')),'*'),'=',sub('z',seq('(',sub('J',seq('m,b')),'+j) mod 60'))),'循环区块抽样',r'r^*_{m,t}=z_{(J_{m,b}+j)\bmod60},\quad j=0,\ldots,4')
eq(seq(sup(sub('S',seq('m,t')),'*'),'=',sub('S','0'),fn('exp',summ('u=1','t',sup(sub('r',seq('m,u')),'*')))),'模拟价格路径',r'S^*_{m,t}=S_0\exp\left(\sum_{u=1}^{t}r^*_{m,u}\right)')
para('每条路径抽取12个区块起点，每块连续读取5个中心化收益；越过样本末尾时循环回头，得到60个未来交易日。不同期限使用同一路径前缀，多阶段保持共同价格因子。循环连接属于研究近似，区块长度尚未完成寻优。')
heading('4 3 候选模型')
eq(seq(sub('r','drift'),'*=z+',mr('0.25'),sub('r','mean')),'历史漂移收缩',r'r^*_{drift}=z_{sample}+0.25\bar r')
eq(seq('a=',fn('min',seq('2,',fn('max',seq('0.5,',frac(sub('σ','20'),sub('σ','60')))))),',  ',sub('r','vol'),'*=az'),'近期波动调整',r'a=\min\{2,\max(0.5,\hat\sigma_{20}/\hat\sigma_{60})\},\quad r^*_{vol}=az_{sample}')
para('25%的收缩和0.5至2的缩放边界是实验设定。基准、漂移及波动三个模型采用相同随机区块起点。它们用于模型比较与企业策略敏感性，尚未证明能够稳定提高预测表现。')

page('5 价格预测与求解步骤')
heading('5 1 阶段末五日均价预测')
table(['交易日','截止日','均值','P10','中位数','P90'],[[x['horizon'],x['expiry'],price(x['average5_mean']),price(x['average5_p10']),price(x['average5_median']),price(x['average5_p90'])] for x in S['forecasts']],[.65,1.4,1.2,1.2,1.2,1.2])
para('单位为元/吨。P10至P90是模型中间80%情景区间。该表是末5日平均价格的分布，保险赔款必须逐条路径先计算，不能把平均价格直接代入非线性赔款函数。')
heading('5 2 期末单日价格预测')
table(['交易日','均值','P10','中位数','P90'],[[x['horizon'],price(x['end_mean']),price(x['end_p10']),price(x['end_median']),price(x['end_p90'])] for x in S['forecasts']],[.85,1.5,1.5,1.5,1.5])
heading('5 3 求解顺序')
steps(['读取连续行情和来源标记，确定起点及交易日历。','构造60个收益率、中心化收益和候选模型参数。','固定种子2026100903，生成三组100,000条60日路径。','按10、20、40、60日截取末5日价格，计算阶段指数。','逐路径计算赔款，再汇总纯保费、尾部资本和总保费。','使用同一价格路径比较企业采购及出售规则，筛选符合预算的方案。'])
para('不能通过读取未来最高价选择买卖日期。20日预测分位较此前略有变化，原因是本版统一了模拟路径与种子；这种变化不能解释为增加了新行情或提升了模型质量。')

page('6 保险责任与价格标准')
eq(seq('K=',frac('1','5'),summ('j=0','4',sub('S','t₀−j'))),'预算价格',r'K=\frac15\sum_{j=0}^{4}S_{t_0-j}')
eq(seq('d=',fn('max',seq('2,0.02K')),',  c=0.05K,  η=0.8'),'候选价格责任参数',r'd=\max(2,0.02K),\quad c=0.05K,\quad\eta=0.8')
eq(seq(sub('X','i'),'=',frac('1','5'),summ('j=0','4',sub('S','Tᵢ−j'))),'阶段结算指数',r'X_i=\frac15\sum_{j=0}^{4}S_{T_i-j}')
eq(seq(sub('L','i'),'=η',sub('Q','i'),fn('min',seq(fn('max',seq(sub('X','i'),'−',sub('K','i'),'−',sub('d','i'),',0')),',',sub('c','i')))),'阶段赔款',r'L_i=\eta Q_i\min\{\max(X_i-K_i-d_i,0),c_i\}')
eq(seq(sub('A','i'),'=η',sub('Q','i'),sub('c','i')),'阶段责任限额',r'A_i=\eta Q_i c_i')
para('K=94.966、d=2、c=4.7483元/吨。触发价格96.966元/吨，达到最高赔款的价格101.7143元/吨。价格超过上沿后仍赔付封顶金额。企业承担免赔、20%共保和超出最高保障价差的成本。')
table(['阶段指数 元每吨','1,000吨赔款 元'],[['95.00','0.00'],['98.00','827.20'],['100.00','2,427.20'],['101.7143','3,798.64'],['105.00','3,798.64']],[3.3,3.55])
para('阶段末均价一次结算，每日收盘价用于组成指数。每天越界不形成对同一吨的重复赔款。公开碳保险条款有预先约定采价窗口及交易日均价的设计可参考；其欧盟期货指数与金额结构应与本方案分别处理。[S3]')

page('7 企业承保数量与金额限额')
eq(seq('D=',fn('max',seq('R−H−C−B,0'))),'未锁价净采购需求',r'D=\max(R-H-C-B,0)')
eq(seq('Q=',fn('min',seq('D,',sub('q','req'),',',sub('q','cap')))),'单次承保数量',r'Q=\min(D,q_{req},q_{cap})')
eq(seq(sub('q','cap,amount'),'=',frac(sub('A','cap'),'ηc')),'金额容量换算',r'q_{cap}^{amount}=A_{cap}/(\eta c),\quad\eta c>0')
para('R为目标年度履约需求；H为合法可用资源；C为未计入H、已锁价且可按期交付的采购；B为未锁价需求中已分配给其他阶段、有效保险或待确认申请的数量。获配与库存去重，同一吨不重复扣减。保险人容量需考虑组合集中度、流动性及资本。')
table(['输入或状态','处理'],[['资料缺失或不清楚','保留未知 列出待补字段 不生成正式额度'],['申请量高于净需求','提示调整申请或核验需求 不静默缩减'],['容量低于申请量','展示容量约束及暂定量 企业确认后推进'],['净需求为零','不生成未来采购保障额度'],['多阶段或多保单','统一企业与年度台账 预占并去重'],['行业类型','选择核算 核验与风险分组资料 不直接推定额度']],[1.65,5.2])
heading('7 1 演示企业额度')
para('演示履约需求100万吨、可用资源98万吨、未计入资源的锁价采购1万吨、已分配需求0吨，净需求为1万吨。申请1万吨且假设批准容量1万吨，暂定承保量1万吨。该示例没有真实企业证据，不代表已批准商业容量。')
table(['承保数量示例','候选最高赔款 元'],[[f'{x["quantity_t"]:,}吨',money(x['limit_cny'])] for x in S['tiers']],[3.3,3.55])
para('最高赔款随Q、c、η共同变化。单次承保数量不限制企业实际采购；每阶段分配量之和仍须满足未锁价净需求约束。')

page('8 纯保费与总保费计算')
eq(seq('μ=',frac('1','M'),summ('m=1','M',sub('L','m')),',  ',sub('Π','pure'),'=',fn('exp',seq('−rτ')),'μ'),'预期赔款与纯保费',r'\hat\mu=M^{-1}\sum_{m=1}^M L_m,\quad\Pi_{pure}=e^{-r\tau}\hat\mu')
eq(seq(sub('TVaR','99.5%'),'=',frac('1','500'),summ('j=M−499','M',sub('L','(j)'))),'模拟尾部赔款',r'\operatorname{TVaR}_{99.5\%}=\frac1{500}\sum_{j=M-499}^{M}L_{(j)},\quad M=100000')
eq(seq('EC=',fn('max',seq(sub('TVaR','99.5%'),'−μ,0')),',  ',sub('C','capital'),'=kτEC'),'资本代理及资本费用',r'EC=\max(\operatorname{TVaR}_{99.5\%}-\hat\mu,0),\quad C_{capital}=k\tau EC')
eq(seq('Π=',frac(seq(sub('Π','pure'),'＋',sub('C','capital'),'＋F＋',sub('C','model')),seq('1−v−p'))),'总保费',r'\Pi=\frac{\Pi_{pure}+C_{capital}+F+C_{model}}{1-v-p}')
table(['参数','研究假设','计量口径'],[['r 折现率','2%','年率 实际日历天数除365'],['k 资本成本','8%','年率 资本代理持有至保障结束'],['v 变动费用','15%','占总保费'],['p 利润','5%','占总保费'],['F 固定费用','200元','每张保单收一次'],['模型附加','基准0% 敏感性5% 10%','占折现纯保费 尚未校准']],[1.6,1.65,3.6])
para('费用率和利润率以总保费为分母，要求v+p<1。TVaR按损失排序的上尾概率质量求均值，封顶产生概率质量时仍按最差0.5%计算；一般路径数需对尾部边界进行分数样本权重处理。')
para('资本代理未按保险人偿付能力口径核定，费用也尚非实际经营报价。赔款、费用、资本和模型不确定性分别估计的原则参考ASOP 53；该文献为方法参考，不是中国保险业务监管标准。[S4]')

page('9 产品定价结果与采价比较')
table(['交易日','日历天','预期赔款','纯保费','资本费用','总保费'],[[x['horizon'],x['calendar_days'],money(x['expected_payout_cny']),money(x['pure_premium_cny']),money(x['capital_cost_cny']),money(x['gross_premium_cny'])] for x in S['quotes']],[.65,.65,1.45,1.45,1.25,1.4])
para('单位元；每档承保1,000吨，最高赔款3,798.64元。单阶段期限延长改变赔款分布与资金持有时间，不能用短期保费按天数线性外推。40、60日尚缺独立赔款检验。')
heading('9 1 20日采价口径比较')
table(['采价方式','预期赔款 元','总保费 元'],[[x['sampling'],money(x['expected_payout_cny']),money(x['gross_premium_cny'])] for x in S['sampling']],[3.1,1.9,1.85])
para('四种方式共用同一价格路径。期末单日适合特定日集中采购；末5日均价适合阶段末分批采购；全期间均价适合贯穿期间的采购；最高收盘价为更宽责任，示例仍只结算一次。选择采价口径应贴合真实采购计划，不能只选择报价最低者。')
heading('9 2 20日保费金额示例')
q=S['quotes'][1]
para(f'预期赔款{money(q["expected_payout_cny"])}元，经35日折现后纯保费{money(q["pure_premium_cny"])}元；资本费用{money(q["capital_cost_cny"])}元、固定费用200元，合计除以0.8，得到总保费{money(q["gross_premium_cny"])}元。')
para(f'赔付概率为{percent(q["probability_claim"])}，预期赔款模拟标准误{money(q["mc_se_expected_payout_cny"])}元。该标准误仅反映随机抽样，不包含模型误差、来源误差或企业执行价差。')
heading('9 3 模型附加敏感性')
table(['附加占纯保费','示范总保费 元'],[[percent(x['margin_rate']),money(x['gross_premium_cny'])] for x in S['margin']],[3.3,3.55])
para('5%和10%仅为敏感性假设，没有证据证明足以覆盖回放中的低估。')

page('10 多阶段主保单定价')
table(['阶段','截止日','吨数','最高赔款 元','单独保费 元'],[[x['horizon'],x['expiry'],x['quantity_t'],money(x['limit_cny']),money(x['gross_premium_cny'])] for x in S['stages']],[.65,1.45,.75,1.85,2.15])
eq(seq('L=',summ('i=1','n',sub('L','i')),',  ',sub('A','total'),'=',summ('i=1','n',sub('A','i'))),'多阶段累计责任',r'L=\sum_i L_i,\quad A_{total}=\sum_i A_i')
eq(seq(sub('Π','pure,joint'),'=',summ('i=1','n',seq(fn('exp',seq('−r',sub('τ','i'))),fn('E',sub('L','i'))))),'多阶段纯保费',r'\Pi_{pure,joint}=\sum_i e^{-r\tau_i}E[L_i]')
eq(seq(sub('C','capital,joint'),'=k',sub('τ','max'),fn('max',seq(sub('TVaR','99.5%'),'(L)−',fn('E','L'),',0'))),'多阶段资本费用',r'C_{capital,joint}=k\tau_{max}\max\{\operatorname{TVaR}_{99.5\%}(L)-E[L],0\}')
para('本例一次安排400、300、300吨，共1,000吨。阶段赔款取自同一价格路径，不能把阶段独立模拟。赔款相关系数约为0.563、0.416、0.692；共用碳价风险意味着不能假设风险完全分散。')
table(['整单指标','结果 元'],[['预期累计赔款',money(J['expected_payout_cny'])],['各阶段折现纯保费之和',money(J['pure_premium_cny'])],['累计最高赔款',money(J['total_limit_cny'])],['资本费用',money(J['capital_cost_cny'])],['主保单示范总保费',money(J['gross_premium_cny'])],['三段分别投保的总保费',money(J['sum_separate_premiums_cny'])]],[3.6,3.25])
para('主保单固定费用仅收200元，三张独立保单各收200元。资本代理保守地持有至最后阶段，未模拟按阶段释放；因此上述差额不是已核实的实际商业折扣。如另设低于各段限额之和的整单上限，须重新截断累计赔款并定价。')

page('11 历史回放与模型问题')
para('20日回放使用当时已知的历史窗口和起点预算价，实际指数取未来第16至20个交易日均价，按相同候选责任计算历史指数赔款。它们是按价格重构的赔款，不是真实保险理赔数据。验证42个起点，后段40个起点，中间清除19个起点。')
eq(seq('RMSE=',root(seq(frac('1','n'),summ('o=1','n',sup(group(seq(sub('μ','o'),'−',sub('L','o'))),'2'))))),'赔款预测误差',r'RMSE=\sqrt{n^{-1}\sum_o(\hat\mu_o-L_o)^2}')
rows=[]
names={'zero60':'区块基准','shrink_drift25':'漂移收缩','recent_vol_scale':'波动调整'}
for model in names:
    v=next(r for r in metrics if r['model']==model and r['partition']=='validation_replay')
    t=next(r for r in metrics if r['model']==model and r['partition']=='test_replay')
    rows.append([names[model],money(v['rmse_mean_payout_cny']),money(t['rmse_mean_payout_cny']),money(t['mean_bias_cny']),f'{float(t["actual_over_expected"]):.3f}'])
table(['模型','验证RMSE','后段RMSE','后段偏差','历史赔款比预期'],rows,[1.3,1.4,1.4,1.4,1.35])
para('按验证RMSE，基准模型为本版研究基准。后段基准预期赔款合计29,060.92元，历史指数赔款合计53,607.44元，后者为前者的1.845倍，说明需要检查趋势、状态变化或分布尾部等低估原因。')
para('两段内部的20日结果窗口重叠；这些日期此前已用于研究。因此42及40个起点不是独立保险风险，本轮也不是新的独立检验。不能把1.845倍直接作为最终加价系数，然后用同一数据声称校准成功。')
heading('11 1 下一轮校准方法')
steps(['预先固定新的验证区间、误差及校准指标。','在每个起点只使用已经成熟的历史赔款标签；20日标签需等待20交易日。','在训练资料内估计漂移、窗口或校准附加，不读取未来验证结果。','用未参与选择的样本评价预期赔款、发生概率与尾部表现。','分别检查10、40、60日和多阶段策略，记录重叠窗口及样本不确定性。'])
para('新报价的100,000条路径不增加独立历史样本。该回放沿用历史种子20261009加起点索引，每起点20,000条路径；本版统一路径种子仅用于当前产品与企业决策，历史回放记录完整保留。')

page('12 降低保费的条件与取舍')
heading('12 1 同路径比较责任变化')
table(['条款','此前责任结构同路径复算','本版候选责任'],[['承保吨数','1,000','1,000'],['预算价 元每吨','94.96','94.966'],['免赔 元每吨','2','2'],['每吨赔付比例','100%','80%'],['最高保障价差 元每吨','10','4.7483'],['最高赔款 元','10,000','3,798.64'],['采价方式','期末单日','末5日均价'],['20日保费 元',money(S['old_contract_same_paths']['gross_premium_cny']),money(S['quotes'][1]['gross_premium_cny'])]],[1.65,2.6,2.6])
para('两列使用同一组价格路径和费用假设。降低保费伴随最高赔款和赔付比例下降、采价口径变化，不能把差额当作算法改进的优惠。此前不同种子得到的1,605.61元等历史报价属于不同复算版本。')
heading('12 2 可实际研究的降费方式')
para('部分投保可以直接减少保障量；免赔和最高价差选择可以调整责任；同企业多阶段合并可摊薄固定费用。费用自动化、再保险与对冲需要实际成本证明，不能事先假设一定节省。')
table(['20日吨数示例','最高赔款 元','总保费 元'],[[f'{x["quantity_t"]:,}',money(x['limit_cny']),money(x['gross_premium_cny'])] for x in S['tiers']],[1.5,2.65,2.7])
para('固定费用按每张保单收取，所以总保费不完全随吨数成比例。大额示例未经过保险人容量及组合风险批准。')
heading('12 3 报价充足性')
para('降低报价之前应先修复赔款预测不足，并计入真实费用及资本。报价优化的验收应分别说明责任变化、费用下降和模型表现变化，不能以保费更低作为算法更好的唯一证据。')

page('13 企业购售决策模型')
eq(seq(sub('I','t'),'=',sub('I','t−1'),'＋',sub('a','t'),'＋',sub('b','t'),'−',sub('s','t'),'−',sub('v','t')),'配额库存守恒',r'I_t=I_{t-1}+a_t+b_t-s_t-v_t')
eq(seq(sub('V','t'),'=',fn('max',seq(sub('I','t'),'−',sub('I','restricted,t'),'−',sub('R','safe,t'),',0'))),'可售余量',r'V_t=\max(I_t-I^{restricted}_t-R^{safe}_t,0)')
eq(seq('0≤',sub('s','t'),'≤',sub('V','t'),',  0≤',sub('Q','ins,i'),'≤',sub('b','unlocked,i')),'交易与保障数量约束',r'0\le s_t\le V_t,\quad0\le Q_{ins,i}\le b_{unlocked,i}')
para('a为到账配额，b为买入并交付量，s为卖出量，v为清缴消耗。R安全包括尚需清缴的预留和经过企业确认的安全缓冲。年度用途与资源合法性需核验，未来合同不作为已交付可售库存。')
eq(seq('θ=',frac(sub('Q','ins'),'D'),',  α=',frac(sub('Q','ins'),seq('D−',sub('q','0')))),'部分投保数量比例',r'\theta=Q_{ins}/D,\quad\alpha=Q_{ins}/(D-q_0)')
para('总缺口1万吨，先买5,000吨，未来采购中保2,500吨，总缺口覆盖率θ=25%，剩余采购覆盖率α=50%；每吨赔付比例η=80%仍单独计算。没有未来未锁价采购时不定义α并不安排该类保险。')
eq(seq(sub('C','net'),'=采购支出＋保费＋交易及资金成本−赔款−出售净回款'),'企业净成本框架',r'C_{net}=\text{采购支出}+\text{保费}+\text{交易及资金成本}-\text{赔款}-\text{出售净回款}')
para('实际买价P实际与指数X可能不同。企业增加或取消采购、改变日期、出售已保障需求对应配额后，应重核保险数量及利益；最终处理以合同为准。采购上涨保险不覆盖售碳价格下跌。')

page('14 企业方案目标与算法')
eq(seq('min  ',fn('E','C'),'＋λ',group(seq(sub('CVaR','95%'),'(C)−',fn('E','C')))),'采购策略风险目标',r'\min E[C]+\lambda(\operatorname{CVaR}_{95\%}(C)-E[C])')
eq(seq(sub('CVaR','95%'),'(C)=',sub('min','z'),group(seq('z＋20',fn('E',sup(group(seq('C−z')),'+'))))),'成本尾部风险定义',r'\operatorname{CVaR}_{95\%}(C)=\min_z\{z+20E[(C-z)^+]\}')
para('λ反映企业对不利成本的权重，演示值为0、0.15和0.5。企业也可以指定现金、保费或成本超预算概率约束。真实风险容忍度需在交互中确认。')
heading('14 1 单阶段成本的时间统一')
eq(seq(sup(sub('C','20'),'(0)'),'=',group(seq(sub('q','0'),sub('S','0'),'(1＋f)＋Π')),fn('exp',seq(sub('r','f'),sub('τ','20')))),'当前支出累积至共同评价时点',r'C^{(0)}_{20}=[q_0S_0(1+f)+\Pi]e^{r_f\tau_{20}}')
eq(seq(sub('C','20'),'=',sup(sub('C','20'),'(0)'),'＋',group(seq(group(seq('D−',sub('q','0'))),sub('X','h'),'(1＋f)−',sub('L','h'))),fn('exp',seq(sub('r','f'),group(seq(sub('τ','20'),'−',sub('τ','h')))))),'未来采购与赔款现金流',r'C_{20}=C^{(0)}_{20}+[(D-q_0)X_h(1+f)-L_h]e^{r_f(\tau_{20}-\tau_h)}')
para('本演示令分批执行价等于末5日指数，全部方案统一到第20交易日计量；提前支出计入资金成本。实际成交价差、赔款到账延迟、税费及融资额度尚未输入。')
heading('14 2 枚举及筛选')
steps(['枚举提前采购比例和剩余采购投保比例各0%、25%、50%、75%、100%，并选10或20交易日采购窗口。','给定同一条款，以基准模型计算保费。没有投保量时保费为零，不收保险固定费。','剔除初期现金、保费预算或容量超额的组合。','在三个价格模型上计算同一方案的平均及尾部成本。','选择最不利候选模型下目标成本最低的可行方案，并展示其他预算可行档位。'])
para('共42种数量与时点组合，三个模型形成126条成本记录。选择仅在候选网格及三个模型内有效；未覆盖全部价格风险，也不是连续变量的全局最优。')

page('15 企业部分投保方案示例')
heading('15 1 企业与费用假设')
table(['假设','演示值'],[['未锁价采购缺口','10,000吨 无可售余量'],['初期采购及保费资金','550,000元'],['保费预算','3,500元'],['假设保险人容量','10,000吨 待实际核验'],['企业年资金成本','6%'],['单边交易费率','0.1% 非实际收费报价'],['采购安排','先买5,000吨 剩余5,000吨固定20日窗口'],['保障条款','末5日均价 免赔2元 赔付80% 价差上限4.7483元']],[1.85,5.0])
cr=[r for r in candidates if r['model']=='zero60' and r['buy_now_t']=='5000.0' and r['future_horizon_trading_days']=='20']
table(['投保吨数','总缺口覆盖','保费 元','平均成本 元','尾部成本 元'],[[f'{float(r["insured_t"]):,.0f}',percent(r['insured_fraction_of_total_gap']),money(r['premium_cny']),money(r['mean_net_cost_cny']),money(r['cvar95_net_cost_cny'])] for r in cr],[1,1.1,1.4,1.65,1.7])
para('尾部成本是基准模型最差5%情景的平均净成本。5,000吨投保方案保费4,276.27元，超过3,500元预算；其他四档可行。投保在本模拟中略提高平均成本，同时降低尾部成本。')
fx=E['fixed20_example']
para(f'采购日期固定20日、λ=0.5时，三个模型的保守比较选择先买5,000吨、后买5,000吨，其中投保3,750吨。总缺口覆盖37.5%，剩余采购覆盖75%，保费{money(fx["premium_cny"])}元，最高赔款14,244.90元。',True)
para('3,750吨不是按行业经验直接指定，而是在演示数量网格、保费预算和风险目标约束下选出。改变费用、采购日期或风险权重时须重新求解。')

page('16 采购时点与压力情景')
table(['偏好或限制','未来窗口','先买吨数','投保吨数','保费 元'],[[x['profile'],x['future_horizon_trading_days'],f'{x["buy_now_t"]:,.0f}',f'{x["insured_t"]:,.0f}',money(x['premium_cny'])] for x in E['buyer_choices'][:2]+[E['fixed20_example']]],[2.35,.75,1.1,1.1,1.55])
para('本组假设中，允许10日窗口采购时，期望成本优先方案不投保；控制尾部成本时，提前买5,000吨，未来5,000吨全部安排保障，覆盖总缺口50%，示范保费2,227.88元。此方案由提前采购和保险共同分担价格风险。')
heading('16 1 总保费变动的影响')
table(['倍率假设','未来窗口','投保吨数','调整后保费 元'],[[x['gross_premium_multiplier_assumption'],x['future_horizon_trading_days'],f'{x["insured_t"]:,.0f}',money(x['adjusted_premium_cny'])] for x in E['gross_premium_sensitivity']],[1.35,1.35,1.6,2.55])
para('倍率是报价敏感性假设，没有用历史1.845倍差异校准。总保费变成2倍时，3,500元预算使候选投保量降为3,750吨。')
heading('16 2 固定20日方案的压力情景')
stress=[r for r in csvrows('采购成本压力情景.csv') if r['plan_id']==fx['plan_id']]
table(['未来采购均价 元每吨','赔款 元','净成本 元'],[[price(r['procurement_index_cny_per_t']),money(r['payout_cny']),money(r['terminal_net_cost_cny'])] for r in stress],[2.75,1.85,2.25])
para('这些均价压力情景不参与原方案的选择。价格超过上沿后赔款封顶，企业净成本仍继续上涨。表内未假设可在压力情景中事后调整采购日期。')

page('17 富余配额出售规则')
heading('17 1 可售余量与利润口径')
para('另一演示企业持有22,000吨，履约及安全缓冲预留20,000吨，无受限库存，可售2,000吨。它与采购缺口企业是两个独立案例，不能将结果相加解释为无风险套利。')
eq(seq(sub('N','τ'),'=',fn('exp',seq('−',sub('r','f'),'τ')),'V',sub('S','τ'),'(1−f)'),'出售净回款现值',r'N_\tau=e^{-r_f\tau}VS_\tau(1-f)')
eq(seq('卖出利润=净出售回款−取得成本−其他未计费用'),'出售利润口径',r'\text{卖出利润}=\text{净出售回款}-\text{取得成本}-\text{其他未计费用}')
para('本轮缺少取得成本及未来自用价值，输出回款不输出利润率。其他未计费用包括尚未扣除的相关费用及税费；已计入净回款的交易费用不再扣一次。资金时间价值已通过折现计入，同一资金成本不应重复扣除。结转、年份及账户限制还需按实际配额核验。')
heading('17 2 比较的出售规则')
steps(['在研究起点出售全部已核验可售余量。','预先约定第10或20交易日出售。','收盘达到98、100或105元后，下一交易日执行；最迟第20日出售。'])
eq(seq('max  ',fn('E','N'),'−λ',group(seq(fn('E','N'),'−',sub('LCVaR','95%'),'(N)'))),'出售下行风险目标',r'\max E[N]-\lambda(E[N]-\operatorname{LCVaR}_{95\%}(N))')
para('LCVaR为最差5%回款情景的均值。选择采用三个模型中最不利的目标回款。阈值只能依据已观察价格触发，实际执行在下一交易日，不能把当天已经结束的收盘价认定为一定可成交。')
para('还可扩展分批出售，例如先变现部分余量，其余按条件出售。本轮完成整批出售规则比较，分批比例、订单簿与实际滑点尚需单独优化。')

page('18 富余配额出售结果')
sr=[r for r in sales if r['model']=='zero60']
labels=['立即出售','第10日出售','第20日出售','98元触发次日执行','100元触发次日执行','105元触发次日执行']
table(['规则','平均回款现值 元','较立即出售 元','更差的模拟比例'],[[label,money(r['mean_pv_net_revenue_cny']),money(r['mean_advantage_vs_now_cny']),percent(r['probability_underperform_now'])] for label,r in zip(labels,sr)],[2.15,1.8,1.5,1.4])
para('基准模型下等待策略没有抵消资金成本，三个模型的保守比较也选择尽早出售。此结论针对9月30日冻结起点、6%资金成本、0.1%单边费用及2,000吨已核验余量，不构成10月9日卖出指令。')
heading('18 1 模型分歧')
ds=next(r for r in sales if r['model']=='shrink_drift25' and r['policy']=='第20交易日出售')
para(f'漂移收缩模型对20日等待的平均优势为{money(ds["mean_advantage_vs_now_cny"])}元，但有{percent(ds["probability_underperform_now"])}的路径回款低于立即出售。价格上涨预期改变会改变等待的评价，不能只凭一个模型确定最佳卖点。')
heading('18 2 企业可接受的输出')
para('智能体应展示可售量、预留依据、出售条件、平均回款、较立即出售差额、下行回款和现金需求。如果真实企业急需现金，当前变现价值可能更高；若未来可能自用配额，出售后的补购机会成本也必须计入。')
para('采购成本上涨保险不覆盖卖价下跌。企业若希望保障售碳收入，需要另行建立下跌赔付责任，采集可售数量及取得成本，再单独测算保费与承保利益。')

page('19 智能体企业交互与输出')
table(['顺序','采集内容','用途'],[['1','企业身份 行业 履约年度 信息时点','选择适用资料与核验规则'],['2','应清缴量 可用资源 已锁价采购','计算净采购需求并去重'],['3','其他阶段及已有保障的分配量','防止同一需求重复保障'],['4','采购数量 日期 交付安排 可否提前','构造可行采购窗口'],['5','现金及保费预算 融资成本 风险偏好','选择投保及采购方案'],['6','可交易库存 受限量 履约预留 缓冲','计算候选可售余量'],['7','取得成本 未来用途 费用 执行价差','计算出售利润与真实净成本'],['8','条款选择 保险人报价 容量 核验状态','生成待核验方案及最终报价']],[.55,3.5,2.8])
para('每个重要字段记录数值、单位、适用年度、信息时点、来源及证据编号。资料未知保持未知；明确为零才作为零计算。预测、企业陈述与已核验数值分别标记。')
heading('19 1 输出给企业的方案字段')
para('提前采购量、未来采购量与窗口、投保吨数、总缺口与剩余采购覆盖率、每吨赔付比例、最高赔款、保费、初期现金、平均及尾部成本、可售量、出售规则、净回款与风险、和基准方案的差额。应同时列出约束原因及尚待核验字段。')
heading('19 2 状态与数量台账')
para('流程依次为资料收集中、研究示例、待核验方案、条款及容量确认、正式报价。统一台账按企业、履约年度和需求批次预占数量，待确认申请也预占；确认时原子占用，取消时释放，资源或采购变化后重新计算。')
para('现有JSON为交互配置，Python为计算原型；正式界面、证据系统与并发台账尚未接入。缺少库存或履约预留时不生成可执行出售量，缺少取得成本时不输出出售利润。')

page('20 期货辅助风险管理')
heading('20 1 期货对保险人的作用')
para('保险赔款按CEA现货指数计算。只有在具备经核验的可交易合约、权限、流动性和风险管理条件时，才能研究用期货降低保险人的价格风险。现有数据没有输入可用于本产品的已核验期货成交、基差、费用或保证金参数，本版对冲收益按零计。')
eq(seq(sub('F','model'),'=',sub('S','0'),fn('exp',seq(group(seq('r＋u−y')),'τ')),'＋b'),'辅助理论期货价格',r'F^{model}_{0,T}=S_0e^{(r+u-y)\tau}+b')
para('该式只能在明确持有成本u、便利收益y、基差调整b及市场约束的假设下提供理论参考。它不是交易所报价，不等于现货价格预测，也不能证明合约可交易或保险可复制。碳配额年份、结转和流动性可能影响该近似。')
eq(seq(sub('G','hedge'),'=q',group(seq(sub('F','T'),'−',sub('F','0')))),'期货对冲收益',r'G_{hedge}=q(F_T-F_0)')
eq(seq(sub('L','residual'),'=L−',sub('G','hedge'),'＋执行及融资成本'),'保险人剩余损失',r'L_{residual}=L-G_{hedge}+\text{执行及融资成本}')
eq(seq('q*=',frac(fn('Cov',seq('L,ΔF')),fn('Var','ΔF'))),'最小方差对冲数量',r'q^*=\operatorname{Cov}(L,\Delta F)/\operatorname{Var}(\Delta F)')
para('保险人赔款随价格上涨增加，正相关时理论对冲方向为多头。实际数量需受交易权限、流动性、保证金和容量限制；封顶赔款与线性期货收益不同，不能保证全额复制。对冲需另测基差、追保流动性和净风险成本。')
heading('20 2 降费前的证据')
para('只有对冲后的赔款与资本成本节省超过交易、基差和融资成本，并通过实际可执行性验证时，才能研究把节省反映到保费。未经验证的理论期货价不用于本版保险加价或折扣。')

page('21 实施顺序与验收')
table(['优先级','当前成果','下一项验收'],[['P0 限额','净需求公式 数量校验 交互字段','真实企业资料去重 容量授权 数量台账'],['P1 价格责任','采价 免赔 比例 价差上限 边界','采购窗口匹配 价格来源 舍入及异常处理'],['P2 保费','共享路径 定价 回放及敏感性','低估校准 独立验证 实际费用与资本'],['P3 企业决策','42组合 6种出售规则 压力情景','执行价差 需求误差 资金及策略回放'],['P4 产品接入','计算原型与交互配置','界面 核验 并发台账 变更及理赔流程']],[1.15,2.6,3.1])
heading('21 1 已完成的计算校验')
para('资料未知、负数量、无效日期、无需求、申请超额和容量错误均有处理。赔款触发及封顶、分阶段数量之和、企业投保量不超过未来采购、次交易日出售触发和产品与企业报价共用路径已检查。计算通过不能代替统计校准或商业承保批准。')
heading('21 2 正式条款仍需确认')
para('固定行情源、交易品种、收盘定义、报价时点、数据更正流程；采价窗口必须在责任期间内且事先约定。价格缺失时按预定异常流程处理，本模拟没有假设任意前向填充。')
para('统一小数位及舍入顺序；本报告计算保留全精度、展示四舍五入，最终舍入规则会影响边界和报价。确认采购变更、数量核验、重复保障、利益核验、赔款到账时间和证据要求。')
heading('21 3 当前报价状态')
para('真实企业输入、最新行情核验、保险人容量及真实费用、新的独立赔款和策略验证仍待完成。本版保留commercial_quote_allowed=false与live_trade_recommendation_allowed=false，研究结果可用于责任选择和模型校准。')

for chunk,start in enumerate(range(0,len(daily),50),1):
    page(f'附录A 2026年日行情 {chunk}')
    para(f'本页列示第{start+1}至{min(start+50,len(daily))}条实际使用的价格。单位元/吨。官表示official_daily，转表示third_party_republication；完整来源网址和质量标记保留在CEA输入日行情.csv。')
    part=daily[start:start+50];rs=[]
    for i in range(0,len(part),2):
        row=[]
        for j in [i,i+1]:
            if j<len(part):row.extend([part[j]['date'],price(part[j]['close_cny_per_t']),'官' if part[j]['source_type']=='official_daily' else '转'])
            else:row.extend(['','',''])
        rs.append(row)
    table(['日期','收盘价','来源','日期','收盘价','来源'],rs,[1.3,1.3,.825,1.3,1.3,.825],size=9.5)
    if chunk==4:para('以上181条用于本版研究。行业排放量与企业配额数据尚未以真实企业输入形式取得，文中100万吨等企业值属于演示假设，不列入行情数据。')

page('附录B 定价与企业方案算法')
heading('B 1 统一情景与产品定价')
code('''读取2026连续收盘价及来源
计算起点前60个对数收益和中心化收益
固定seed = 2026100903  M = 100000
对三个候选模型共用区块起点 生成60日路径
对h ∈ {10,20,40,60}
    X = 路径第h-4至h日价格均值
    L = η × Q × clip(X-K-d, 0, c)
    pure = exp(-r×日历天数/365) × mean(L)
    EC = max(最差0.5%赔款均值-mean(L), 0)
    capital = k × 日历天数/365 × EC
    premium = (pure+capital+200+模型附加)/0.8
主保单使用同一路径各阶段赔款之和
纯保费分阶段折现 固定费一次 资本持有至末阶段''')
heading('B 2 企业采购与部分投保')
code('''对未来窗口h ∈ {10,20}
    对提前买入比例β ∈ {0,0.25,0.5,0.75,1}
        对剩余采购投保比例α ∈ {0,0.25,0.5,0.75,1}
            q0 = βD  Qins = α(D-q0)
            计算保费 初期现金与容量约束
            对三个模型计算同一方案的净成本
            目标 = E[C]+λ(CVaR95(C)-E[C])
选择最不利模型下目标最低的预算可行组合
保留其他可行档位 展示约束和未核验字段''')
para('程序入口prepare_cea_complete_plan_20261009.py；依赖Python、NumPy和包内辅助文件。重算会更新data目录及统一计算结果.json。图文主文件由build_cea_complete_plan_20261009.py生成，Word公式为原生Office数学格式。')

page('附录C 数据文件与来源')
heading('C 1 主要数据文件')
table(['文件组','内容'],[['CEA输入日行情.csv','原始价格 日期 来源网址 质量标记'],['统一行情预测及单期限保费','共享路径预测 分位 纯保费 资本及总保费'],['统一分阶段保障及计算结果JSON','阶段数量 最高赔款 联合风险与主保单'],['历史赔款预测记录及回放指标','每个起点的预测 价格重构赔款 算法比较'],['全部采购投保候选及演示采购方案','42方案 三模型成本 预算可行性'],['富余配额出售比较及演示出售方案','出售规则 净回款 下行风险'],['保费变动及采购压力情景','费用假设变化 价格85至115元情景'],['智能体输入流程及空白输入示例','字段含义 待补资料 台账及正式报价要求']],[2.95,3.9])
para('原行情文件SHA256为 '+S['input_sha256']+'。包内manifest.json列出实际交付文件的SHA256，可用于核对版本与复算资料。')
heading('C 2 外部依据')
sources=[('S1','生态环境部 碳排放权交易管理暂行条例','履约与配额购售 第十四条','https://www.mee.gov.cn/zcwj/gwywj/202402/t20240205_1065850.shtml'),
('S2','上海环境能源交易所 碳排放权交易管理规则试行','全国市场可交易量及资金约束','https://overview.cneeex.com/c/2021-05-17/491282.shtml'),
('S3','太保 航运业欧盟碳排放成本价格指数保险2025版条款','参考采价窗口及交易日均价设计','https://www.cpic.com.cn/upload/resources/file/2026/04/14/94231.pdf'),
('S4','Actuarial Standards Board ASOP 53','未来风险成本 费用及资本估计的方法参考','https://www.actuarialstandardsboard.org/wp-content/uploads/2018/01/asop053_190.pdf')]
for no,title,use,url in sources:
    para(f'[{no}] {title}。用途：{use}。核阅日期2026年10月9日。\n{url}')
para('S3的欧盟期货指数保险与本CEA现货方案的条款及金额定义不同；S4不是中国监管标准。行情逐条来源以数据文件为准。')

out=OUT/'CEA现货采购保险及企业购售决策完整方案.docx'
doc.core_properties.title='CEA现货采购保险及企业购售决策完整方案'
doc.core_properties.subject='现货模型 保险定价与限额 企业采购投保 富余配额出售'
doc.core_properties.author='碳保险研究项目'
doc.save(out)
(OUT/'完整方案.md').write_text('\n'.join(md),encoding='utf-8')
(OUT/'公式索引.json').write_text(json.dumps(equations,ensure_ascii=False,indent=2),encoding='utf-8')
with zipfile.ZipFile(out) as f:
    x=f.read('word/document.xml').decode()
    assert x.count('<m:oMath>')==len(equations) and '<w:pBdr' not in x
    assert all(needle in x for needle in ['1,055.25','3,269.70','181','1.845'])
(OUT/'文档结构校验.json').write_text(json.dumps(dict(native_equations=len(equations),tables=len(doc.tables),
    daily_rows=len(daily),data_consistency='passed',render_status='pending'),ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(dict(docx=str(out.resolve()),native_equations=len(equations),tables=len(doc.tables),daily_rows=len(daily)),ensure_ascii=False))
